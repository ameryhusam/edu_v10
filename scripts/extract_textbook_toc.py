#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Edu7 Textbook TOC Extractor and Importer Generator
--------------------------------------------------
أداة بايثون متخصصة لاستخراج فهرس كتب PDF المدرسية (الوحدات والدروس وأرقام الصفحات)
من الصفحات العشر الأولى، مع دعم كامل لتخطيط الصفحات ذات العمودين (Two-Column Layouts)،
وتوليد ملف JSON مطابق لمعيار Edu7 Content Package v1.2 وملف Excel متوافق مع القالب المعتمد.

المتطلبات:
    pip install pypdf openpyxl

الاستخدام:
    python scripts/extract_textbook_toc.py <path_to_pdf> [options]
    
أمثلة:
    python scripts/extract_textbook_toc.py "d:/books/math_g07_t1.pdf"
    python scripts/extract_textbook_toc.py "science_g07.pdf" --subject SCI --grade G07 --term 1 --output-dir ./out
"""

import os
import re
import sys
import json
import argparse
from typing import List, Dict, Any, Optional, Tuple

try:
    import pypdf
except ImportError:
    print("خطأ: مكتبة pypdf غير مثبتة. يرجى تثبيتها عبر: pip install pypdf")
    sys.exit(1)

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


# ─── ثوابت ودوال مساعدة للأرقام واللغة العربية ──────────────────────────────

ARABIC_INDIC_DIGITS = "٠١٢٣٤٥٦٧٨٩"
WESTERN_DIGITS = "0123456789"
DIGIT_TRANS = str.maketrans(ARABIC_INDIC_DIGITS, WESTERN_DIGITS)

def to_western_digits(text: str) -> str:
    """تحويل الأرقام العربية المشرقية (٠-٩) إلى أرقام غربية (0-9)."""
    return text.translate(DIGIT_TRANS)

ARABIC_ORDINAL_MAP = {
    "الأولى": 1, "الأول": 1,
    "الثانية": 2, "الثاني": 2,
    "الثالثة": 3, "الثالث": 3,
    "الرابعة": 4, "الرابع": 4,
    "الخامسة": 5, "الخامس": 5,
    "السادسة": 6, "السادس": 6,
    "السابعة": 7, "السابع": 7,
    "الثامنة": 8, "الثامن": 8,
    "التاسعة": 9, "التاسع": 9,
    "العاشرة": 10, "العاشر": 10,
    "الحادية عشرة": 11, "الحادي عشر": 11,
    "الثانية عشرة": 12, "الثاني عشر": 12,
}

TOC_KEYWORDS = [
    "المحتويات",
    "فهرس المحتويات",
    "فهرس الموضوعات",
    "الفهرس",
    "الموضوع",
    "الصفحة",
    "الوحدة الأولى",
    "الوحدة 1",
    "الوحدة الاولي",
]


# ─── استخراج النصوص مع دعم العمودين (Two-Column Extractor) ───────────────────

def is_page_two_columns(page: pypdf.PageObject) -> bool:
    """
    فحص ما إذا كانت الصفحة مقسمة إلى عمودين عبر فحص إحداثيات النصوص الأفقية (X-coordinates).
    """
    mb = page.mediabox
    width = float(mb.right - mb.left)
    mid_x = float(mb.left) + (width / 2.0)
    
    left_count = 0
    right_count = 0
    
    def visitor(text, cm, tm, font_dict, font_size):
        nonlocal left_count, right_count
        if text.strip():
            x = tm[4]
            if x < mid_x - (width * 0.05):
                left_count += 1
            elif x > mid_x + (width * 0.05):
                right_count += 1
                
    try:
        page.extract_text(visitor_text=visitor)
    except Exception:
        return False
        
    total = left_count + right_count
    if total > 20 and left_count > 8 and right_count > 8:
        return True
    return False


def extract_page_text_two_columns(page: pypdf.PageObject) -> str:
    """
    استخراج نصوص الصفحة المقسمة لعمودين مع مراعاة اتجاه القراءة العربي (RTL):
    يتم استخراج نصوص العمود الأيمن أولاً (من الأعلى للأسفل)،
    ثم نصوص العمود الأيسر ثانياً (من الأعلى للأسفل).
    """
    mb = page.mediabox
    width = float(mb.right - mb.left)
    mid_x = float(mb.left) + (width / 2.0)
    
    right_chunks: List[Tuple[float, float, str]] = []
    left_chunks: List[Tuple[float, float, str]] = []
    
    def visitor(text, cm, tm, font_dict, font_size):
        clean = text.strip()
        if not clean:
            return
        x = tm[4]
        y = tm[5]
        if x >= mid_x:
            right_chunks.append((y, -x, clean))
        else:
            left_chunks.append((y, -x, clean))
            
    try:
        page.extract_text(visitor_text=visitor)
    except Exception:
        return page.extract_text()
        
    right_chunks.sort(key=lambda item: (-item[0], item[1]))
    left_chunks.sort(key=lambda item: (-item[0], item[1]))
    
    lines: List[str] = []
    
    # تجميع العمود الأيمن
    curr_y = None
    curr_line: List[str] = []
    for y, neg_x, txt in right_chunks:
        if curr_y is None or abs(y - curr_y) < 6.0:
            curr_line.append(txt)
            curr_y = y
        else:
            lines.append(" ".join(curr_line))
            curr_line = [txt]
            curr_y = y
    if curr_line:
        lines.append(" ".join(curr_line))
        
    # تجميع العمود الأيسر
    curr_y = None
    curr_line = []
    for y, neg_x, txt in left_chunks:
        if curr_y is None or abs(y - curr_y) < 6.0:
            curr_line.append(txt)
            curr_y = y
        else:
            lines.append(" ".join(curr_line))
            curr_line = [txt]
            curr_y = y
    if curr_line:
        lines.append(" ".join(curr_line))
        
    return "\n".join(lines)


# ─── محرك فحص واكتشاف صفحات الفهرس (TOC Detector) ───────────────────────────

def find_toc_pages(reader: pypdf.PdfReader, max_search_pages: int = 10) -> List[Tuple[int, str]]:
    """
    البحث في الصفحات الأولى (افتراضياً حتى 10) للعثور على صفحات الفهرس واستخراج نصوصها.
    """
    toc_pages: List[Tuple[int, str]] = []
    total_pages = len(reader.pages)
    search_limit = min(max_search_pages, total_pages)
    
    for page_idx in range(search_limit):
        page = reader.pages[page_idx]
        raw_sample = page.extract_text() or ""
        normalized_sample = to_western_digits(raw_sample)
        
        is_toc = False
        first_few_lines = "\n".join(normalized_sample.splitlines()[:8])
        for kw in TOC_KEYWORDS:
            if kw in first_few_lines or kw in normalized_sample[:300]:
                is_toc = True
                break
                
        if not is_toc and "الوحدة" in first_few_lines:
            if re.search(r'(\.|\-|…|\s){3,}\s*\d+', normalized_sample):
                is_toc = True
                
        if is_toc:
            if is_page_two_columns(page):
                extracted = extract_page_text_two_columns(page)
            else:
                extracted = page.extract_text() or ""
            toc_pages.append((page_idx + 1, extracted))
            
    return toc_pages


# ─── محلل بنية الفهرس (TOC Hierarchy Parser) ────────────────────────────────

class TocParser:
    def __init__(self, raw_text: str):
        self.raw_text = to_western_digits(raw_text)
        self.units: List[Dict[str, Any]] = []

    def clean_line(self, line: str) -> str:
        line = re.sub(r'[\.]{2,}', ' ', line)
        line = re.sub(r'[-]{3,}', ' ', line)
        line = re.sub(r'[ـ]{2,}', ' ', line)
        line = re.sub(r'[…]+', ' ', line)
        line = re.sub(r'\s+', ' ', line).strip()
        return line

    def extract_page_number_from_line(self, line: str) -> Tuple[str, Optional[int]]:
        line = line.strip()
        m_end = re.search(r'(?:ص|صفحة)?\s*(\d+)\s*$', line)
        if m_end:
            page_num = int(m_end.group(1))
            cleaned_title = line[:m_end.start()].strip()
            return cleaned_title, page_num
            
        m_start = re.search(r'^\s*(\d+)\s*(?:ص|صفحة)?\s+', line)
        if m_start:
            page_num = int(m_start.group(1))
            cleaned_title = line[m_start.end():].strip()
            return cleaned_title, page_num
            
        return line, None

    def parse(self) -> List[Dict[str, Any]]:
        lines = [self.clean_line(l) for l in self.raw_text.splitlines() if l.strip()]
        
        current_unit: Optional[Dict[str, Any]] = None
        unit_counter = 0
        lesson_counter = 0

        for line in lines:
            if any(h in line for h in ["المحتويات", "فهرس المحتويات", "الموضوع / الصفحة", "الموضوع", "فهرس الموضوعات"]):
                if len(line.split()) <= 4:
                    continue

            is_unit_header = bool(re.search(r'^(?:الوحدة|الفصل|الباب)\s+(?:الأولى|الثانية|الثالثة|الرابعة|الخامسة|السادسة|السابعة|الثامنة|التاسعة|العاشرة|\d+)', line))
            
            if is_unit_header:
                unit_counter += 1
                lesson_counter = 0
                title, page_num = self.extract_page_number_from_line(line)
                
                ord_match = re.search(r'(الأولى|الثانية|الثالثة|الرابعة|الخامسة|السادسة|السابعة|الثامنة|التاسعة|العاشرة|\d+)', title)
                ord_val = unit_counter
                if ord_match:
                    w = ord_match.group(1)
                    ord_val = int(w) if w.isdigit() else ARABIC_ORDINAL_MAP.get(w, unit_counter)

                current_unit = {
                    "orderIndex": ord_val,
                    "slug": f"U{ord_val:02d}",
                    "name": title,
                    "startPage": page_num or 1,
                    "endPage": page_num or 1,
                    "lessons": []
                }
                self.units.append(current_unit)
                continue

            title, page_num = self.extract_page_number_from_line(line)

            if page_num is not None:
                if current_unit is None:
                    unit_counter = 1
                    current_unit = {
                        "orderIndex": 1,
                        "slug": "U01",
                        "name": "الوحدة الأولى: عام",
                        "startPage": page_num,
                        "endPage": page_num,
                        "lessons": []
                    }
                    self.units.append(current_unit)

                lesson_counter += 1
                clean_title = re.sub(r'^(?:الدرس|الموضوع)\s*(?:الأول|الثاني|الثالث|الرابع|الخامس|\d+)?[:\-\.]?\s*', '', title).strip()
                if not clean_title:
                    clean_title = title

                lesson_obj = {
                    "orderIndex": lesson_counter,
                    "slug": f"L{lesson_counter:02d}",
                    "name": title,
                    "startPage": page_num,
                    "endPage": page_num,
                    "concepts": [
                        {
                            "orderIndex": 1,
                            "slug": "C01",
                            "name": clean_title,
                            "difficulty": 0.3,
                            "masteryThreshold": 0.8,
                            "isCore": True
                        }
                    ]
                }
                current_unit["lessons"].append(lesson_obj)

        self._calculate_end_pages()
        return self.units

    def _calculate_end_pages(self):
        for u_idx, unit in enumerate(self.units):
            lessons = unit["lessons"]
            for l_idx, lesson in enumerate(lessons):
                if l_idx + 1 < len(lessons):
                    next_start = lessons[l_idx + 1]["startPage"]
                    lesson["endPage"] = max(lesson["startPage"], next_start - 1)
                elif u_idx + 1 < len(self.units) and self.units[u_idx + 1]["lessons"]:
                    next_unit_start = self.units[u_idx + 1]["lessons"][0]["startPage"]
                    lesson["endPage"] = max(lesson["startPage"], next_unit_start - 1)
                else:
                    lesson["endPage"] = lesson["startPage"] + 8

            if lessons:
                unit["startPage"] = lessons[0]["startPage"]
                unit["endPage"] = lessons[-1]["endPage"]
            elif u_idx + 1 < len(self.units):
                unit["endPage"] = max(unit["startPage"], self.units[u_idx + 1]["startPage"] - 1)
            else:
                unit["endPage"] = unit["startPage"] + 20


# ─── توليد مخرجات JSON و Excel المتوافقة مع Edu7 ────────────────────────────

def generate_edu7_json(
    units: List[Dict[str, Any]],
    textbook_title: str,
    subject_key: str,
    grade_key: str,
    part: str,
    edition: str = "2026",
    output_path: str = "curriculum.json"
) -> Dict[str, Any]:
    pad_grade = grade_key.replace("G", "").zfill(2)
    part_code = {"PART_1": "P1", "PART_2": "P2", "BOTH": "PB"}.get(str(part).upper())
    if not part_code:
        raise ValueError("Physical part must be PART_1, PART_2, or BOTH")
    textbook_key = f"EDU-{subject_key}-G{pad_grade}-{part_code}-ED{edition}"
    
    total_pages = 0
    if units and units[-1].get("endPage"):
        total_pages = units[-1]["endPage"]

    flat_units = []
    flat_lessons = []
    flat_concepts = []

    for u in units:
        flat_units.append({
            "slug": u["slug"],
            "parentUnitSlug": None,
            "name": u["name"],
            "orderIndex": u["orderIndex"],
            "startPage": u["startPage"],
            "endPage": u["endPage"],
            "sourceRef": f"{textbook_key}-{u['slug']}"
        })
        for l in u.get("lessons", []):
            flat_lessons.append({
                "unitSlug": u["slug"],
                "slug": l["slug"],
                "name": l["name"],
                "description": f"شرح مفردات درس: {l['name']}",
                "orderIndex": l["orderIndex"],
                "estimatedMins": 45,
                "startPage": l["startPage"],
                "endPage": l["endPage"],
                "sourceRef": f"{textbook_key}-{u['slug']}-{l['slug']}"
            })
            for c in l.get("concepts", []):
                flat_concepts.append({
                    "unitSlug": u["slug"],
                    "lessonSlug": l["slug"],
                    "slug": c["slug"],
                    "name": c["name"],
                    "description": f"المفهوم التعليمي لـ {c['name']}",
                    "orderIndex": c["orderIndex"],
                    "difficulty": c.get("difficulty", 0.3),
                    "importance": 0.8,
                    "masteryThreshold": c.get("masteryThreshold", 0.8),
                    "isCore": c.get("isCore", True),
                    "pageNumber": l["startPage"],
                    "sourceRef": f"{textbook_key}-{u['slug']}-{l['slug']}-{c['slug']}"
                })

    package = {
        "meta": {
            "profile": "edu7.textbook-content",
            "profileVersion": "1.2",
            "scope": "FULL",
            "exportedAt": "2026-09-18T20:00:00.000Z"
        },
        "textbook": {
            "key": textbook_key,
            "subjectKey": subject_key,
            "gradeKey": grade_key,
            "part": part,
            "title": textbook_title,
            "edition": edition,
            "description": f"منهج {textbook_title} المعتمد",
            "status": "DRAFT",
            "publishYear": int(edition) if edition.isdigit() else 2026,
            "totalPages": total_pages
        },
        "units": flat_units,
        "lessons": flat_lessons,
        "concepts": flat_concepts,
        "prerequisites": [],
        "misconceptions": [],
        "learningResources": [],
        "flashcards": [],
        "questions": []
    }

    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(package, f, ensure_ascii=False, indent=2)

    print(f"✅ تم إنشاء ملف JSON بنجاح: {output_path}")
    return package


def generate_edu7_excel(
    units: List[Dict[str, Any]],
    output_path: str = "curriculum.xlsx"
):
    if not HAS_OPENPYXL:
        print("⚠️ مكتبة openpyxl غير مثبتة، تم تخطي إنشاء ملف الإكسل. لتثبيتها: pip install openpyxl")
        return

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "01_الهيكل_والمفاهيم"
    ws.sheet_view.rightToLeft = True

    headers = [
        "رقم الوحدة",
        "اسم الوحدة",
        "رقم الدرس",
        "اسم الدرس",
        "من صفحة",
        "إلى صفحة",
        "اسم المفهوم",
        "عتبة الإتقان",
        "صعوبة المفهوم"
    ]

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Tajawal", size=11, bold=True, color="FFFFFF")
    cell_font = Font(name="Tajawal", size=10)
    center_align = Alignment(horizontal="center", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")

    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    ws.append(headers)
    for col_num, cell in enumerate(ws[1], 1):
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    row_idx = 2
    for u in units:
        u_num = u["orderIndex"]
        u_name = u["name"]
        for l in u.get("lessons", []):
            l_num = l["orderIndex"]
            l_name = l["name"]
            start_p = l["startPage"]
            end_p = l["endPage"]
            for c in l.get("concepts", []):
                c_name = c["name"]
                ws.append([
                    u_num,
                    u_name,
                    l_num,
                    l_name,
                    start_p,
                    end_p,
                    c_name,
                    0.80,
                    3
                ])
                for col in range(1, 10):
                    c_cell = ws.cell(row=row_idx, column=col)
                    c_cell.font = cell_font
                    c_cell.border = thin_border
                    if col in [1, 3, 5, 6, 8, 9]:
                        c_cell.alignment = center_align
                    else:
                        c_cell.alignment = right_align
                row_idx += 1

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 5, 14)

    wb.save(output_path)
    print(f"✅ تم إنشاء ملف Excel بنجاح: {output_path}")


# ─── واجهة سطر الأوامر (CLI Main) ───────────────────────────────────────────

def infer_metadata_from_filename(filename: str) -> Tuple[str, str, str, str]:
    base = os.path.basename(filename).upper()
    
    subject = "GENERAL"
    if "MATH" in base or "رياضيات" in base: subject = "MATH"
    elif "SCI" in base or "علوم" in base: subject = "SCI"
    elif "ARAB" in base or "عربي" in base: subject = "ARAB"
    elif "ENG" in base or "انجليزي" in base: subject = "ENG"
    elif "ISLAM" in base or "اسلامية" in base: subject = "ISLAM"

    grade = "G07"
    g_match = re.search(r'G(\d{1,2})', base)
    if g_match:
        grade = f"G{int(g_match.group(1)):02d}"
    elif "سابع" in base: grade = "G07"
    elif "ثامن" in base: grade = "G08"
    elif "تاسع" in base: grade = "G09"

    part = "PART_1"
    if "T2" in base or "P2" in base or "الفصل_الثاني" in base or "ترم2" in base:
        part = "PART_2"

    title = f"كتاب {subject} للصف {grade} الجزء {1 if part == 'PART_1' else 2}"
    return subject, grade, part, title


def main():
    parser = argparse.ArgumentParser(
        description="استخراج فهرس كتاب مدرسي PDF وتحويله لـ JSON و Excel متوافق مع منصة Edu7"
    )
    parser.add_argument("pdf_path", help="مسار ملف كتاب الـ PDF")
    parser.add_argument("--output-dir", "-o", default="./extracted_curriculum", help="مجلد حفظ الملفات المستخرجة")
    parser.add_argument("--subject", help="رمز المادة (مثال: MATH, SCI)")
    parser.add_argument("--grade", help="رمز الصف (مثال: G07, G08)")
    parser.add_argument("--part", choices=["PART_1", "PART_2", "BOTH"], help="الجزء الفيزيائي للكتاب")
    parser.add_argument("--title", help="عنوان الكتاب المعتمد")
    parser.add_argument("--edition", default="2026", help="طبعة الكتاب (افتراضي: 2026)")
    parser.add_argument("--max-pages", type=int, default=10, help="أقصى عدد صفحات للبحث عن الفهرس من البداية (افتراضي: 10)")
    parser.add_argument("--two-column", action="store_true", help="فرض معالجة تخطيط العمودين")

    args = parser.parse_args()

    if not os.path.exists(args.pdf_path):
        print(f"❌ خطأ: الملف غير موجود: {args.pdf_path}")
        sys.exit(1)

    os.makedirs(args.output_dir, exist_ok=True)

    print(f"🔍 بدء فحص الكتاب: {args.pdf_path}")
    reader = pypdf.PdfReader(args.pdf_path)
    total_pages = len(reader.pages)
    print(f"📄 إجمالي صفحات الكتاب: {total_pages}")

    print(f"🔎 البحث عن صفحات الفهرس في أول {args.max_pages} صفحات...")
    toc_pages = find_toc_pages(reader, max_search_pages=args.max_pages)

    if not toc_pages:
        print("⚠️ لم يتم التعرف التلقائي على صفحة الفهرس بالكلمات المفتاحية المعتادة.")
        print("سحب أول 5 صفحات للفحص المباشر...")
        for p in range(min(5, total_pages)):
            page = reader.pages[p]
            text = extract_page_text_two_columns(page) if args.two_column else (page.extract_text() or "")
            toc_pages.append((p + 1, text))

    print(f"📋 تم العثور على {len(toc_pages)} صفحة/صفحات تمثل الفهرس: {[p[0] for p in toc_pages]}")
    combined_toc_text = "\n".join(p[1] for p in toc_pages)

    print("⚙️ تحليل البنية الهرمية للوحدات والدروس...")
    parser_engine = TocParser(combined_toc_text)
    units = parser_engine.parse()

    if not units:
        print("❌ لم يتمكن المحلل من استخراج أي وحدات أو دروس من النصوص المستخرجة.")
        print("عينة من النص المستخرج:")
        print(combined_toc_text[:500])
        sys.exit(1)

    print(f"✨ تم استخراج بنجاح: {len(units)} وحدة، وإجمالي {sum(len(u.get('lessons', [])) for u in units)} درساً.")
    for u in units:
        print(f"  • {u['name']} (ص {u['startPage']} - {u['endPage']}) — {len(u['lessons'])} درساً")

    inf_subj, inf_grade, inf_part, inf_title = infer_metadata_from_filename(args.pdf_path)
    subject = args.subject or inf_subj
    grade = args.grade or inf_grade
    part = args.part or inf_part
    title = args.title or inf_title

    base_name = f"{grade}-{subject}-{part.replace("PART_", "P")}"
    json_path = os.path.join(args.output_dir, f"{base_name}.json")
    xlsx_path = os.path.join(args.output_dir, f"{base_name}.xlsx")

    generate_edu7_json(
        units=units,
        textbook_title=title,
        subject_key=subject,
        grade_key=grade,
        part=part,
        edition=args.edition,
        output_path=json_path
    )

    generate_edu7_excel(
        units=units,
        output_path=xlsx_path
    )

    print("\n🎉 تمت العملية بنجاح! الملفات جاهزة للاستيراد المباشر في لوحة إدارة Edu7.")

if __name__ == "__main__":
    main()
